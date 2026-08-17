from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "PRD-Crypto-Exist.docx"

NAVY = "101828"
BLUE = "175CD3"
CYAN = "0E7490"
MUTED = "667085"
LIGHT = "F2F4F7"
PALE_BLUE = "EFF8FF"
PALE_CYAN = "ECFDFF"
WHITE = "FFFFFF"
RED = "B42318"
GREEN = "027A48"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=None, bold=None, color=NAVY, italic=None, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.extend([begin, instr, separate, text, end])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    tokens = {
        "Title": (28, NAVY, 0, 6),
        "Subtitle": (13, MUTED, 0, 10),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (11.5, CYAN, 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1", "Heading 2") else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name.startswith("Heading") or name == "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("CRYPTO EXIST  |  PRODUCT REQUIREMENTS DOCUMENT")
    set_font(run, size=8, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("Crypto Exist  |  Internal Product Document  |  ")
    set_font(run, size=8, color=MUTED)
    add_field(fp, "PAGE")


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.75)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)
    r = p.add_run(f"{label}: ")
    set_font(r, size=10.5, bold=True, color=color)
    r = p.add_run(text)
    set_font(r, size=10.5, color=NAVY)


def add_table(doc, headers, rows, widths, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=9 if compact else 9.5, bold=True, color=NAVY)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=8.7 if compact else 9.2, color=NAVY)
    return table


def add_section(doc, title, level=1):
    doc.add_heading(title, level=level)


def build():
    doc = Document()
    configure_document(doc)

    # Cover / memo masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_font(r, size=10, bold=True, color=CYAN)

    p = doc.add_paragraph(style="Title")
    p.add_run("Crypto Exist")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Platform Media, Berita, dan Edukasi Web3 & Blockchain")

    add_table(doc, ["Atribut", "Detail"], [
        ("Status", "Draft awal untuk validasi produk"),
        ("Versi", "1.0"),
        ("Tanggal", date.today().strftime("%d %B %Y")),
        ("Pemilik Produk", "Crypto Exist"),
        ("Audiens", "Founder, product, design, engineering, editorial, dan stakeholder"),
        ("Platform", "Website responsif, mobile-first"),
    ], [2400, 6960])

    add_callout(doc, "Ringkasan", "Crypto Exist akan menjadi media independen Indonesia yang mengutamakan berita kripto yang cepat, edukasi yang mudah dipahami, analisis yang transparan, dan pengalaman editorial yang aman serta efisien.", PALE_CYAN, CYAN)

    doc.add_paragraph("Dokumen ini mendefinisikan kebutuhan produk tingkat MVP. Keputusan teknologi dan layanan pihak ketiga yang belum dikonfirmasi tetap dianggap provisional.")

    doc.add_page_break()
    add_section(doc, "Daftar Isi", 1)
    toc_items = [
        "1. Konteks dan peluang", "2. Visi, tujuan, dan prinsip", "3. Audiens dan persona",
        "4. Ruang lingkup produk", "5. Information architecture", "6. Kebutuhan fungsional",
        "7. Alur pengguna utama", "8. Sistem editorial", "9. Model data konseptual",
        "10. UX dan design system", "11. SEO dan distribusi", "12. Kredibilitas dan tata kelola",
        "13. Keamanan dan privasi", "14. Persyaratan nonfungsional", "15. Analytics dan metrik",
        "16. Roadmap", "17. Acceptance criteria", "18. Risiko dan mitigasi",
        "19. Dependency dan keputusan terbuka", "20. Definition of Done dan rilis"
    ]
    for item in toc_items:
        add_bullet(doc, item)
    doc.add_paragraph("Nomor halaman dapat diperbarui otomatis di Microsoft Word setelah dokumen berkembang.").italic = True

    doc.add_page_break()
    add_section(doc, "1. Konteks dan Peluang", 1)
    doc.add_paragraph("Ekosistem aset kripto dan Web3 bergerak cepat, sementara pembaca Indonesia menghadapi informasi yang terfragmentasi, jargon teknis, konten promosi terselubung, dan analisis tanpa sumber yang memadai. Crypto Exist hadir sebagai media independen yang menggabungkan kecepatan berita, kedalaman edukasi, dan disiplin editorial.")
    add_section(doc, "1.1 Pernyataan masalah", 2)
    for text in [
        "Pembaca sulit membedakan berita, opini, analisis, siaran pers, dan konten sponsor.",
        "Informasi pasar cepat berubah dan sering tidak menyebut sumber maupun waktu pengambilan data.",
        "Materi edukasi sering terlalu teknis untuk pemula atau terlalu dangkal untuk pembaca berpengalaman.",
        "Tim redaksi membutuhkan workflow yang menjaga kecepatan tanpa mengorbankan review dan jejak koreksi.",
    ]:
        add_bullet(doc, text)
    add_section(doc, "1.2 Peluang", 2)
    doc.add_paragraph("Membangun destination media berbahasa Indonesia yang kredibel, cepat, mobile-first, dan memiliki tata kelola editorial eksplisit. Diferensiasi utama bukan hanya volume konten, melainkan transparansi sumber, kualitas penjelasan, dan pengalaman membaca yang baik.")

    add_section(doc, "2. Visi, Tujuan, dan Prinsip", 1)
    add_callout(doc, "Visi", "Menjadi referensi tepercaya bagi masyarakat Indonesia untuk memahami dan mengikuti perkembangan aset kripto, blockchain, Web3, serta DeFi.")
    add_section(doc, "2.1 Tujuan MVP", 2)
    goals = [
        ("G1", "Menyediakan pengalaman membaca berita dan edukasi yang cepat, nyaman, serta mudah ditemukan."),
        ("G2", "Menyediakan workflow redaksi dari draft sampai publikasi dengan role dan audit yang jelas."),
        ("G3", "Membangun fondasi SEO, distribusi, keamanan, dan observability yang layak produksi."),
        ("G4", "Menciptakan mekanisme kepercayaan: sumber, disclosure, penulis, pembaruan, dan koreksi."),
    ]
    add_table(doc, ["ID", "Tujuan"], goals, [900, 8460])
    add_section(doc, "2.2 Non-goals MVP", 2)
    for text in ["Trading atau custody aset", "Koneksi wallet", "Forum dan komentar publik", "Nasihat investasi personal", "Paywall", "Publikasi otomatis tanpa persetujuan manusia"]:
        add_bullet(doc, text)
    add_section(doc, "2.3 Prinsip keputusan", 2)
    for text in ["Kredibilitas sebelum kecepatan", "Security dan privacy by design", "Mobile-first", "Accessible by default", "Server-first dan progressive enhancement", "Sederhana, reversibel, dan terukur"]:
        add_bullet(doc, text)

    add_section(doc, "3. Audiens dan Persona", 1)
    personas = [
        ("Pemula kripto", "Memahami istilah, risiko, wallet, dan dasar blockchain", "Jargon, hoaks, rasa takut salah", "Edukasi bertahap, glosarium, disclaimer"),
        ("Investor aktif", "Mengikuti berita dan konteks pasar dengan cepat", "Informasi terlambat atau tanpa sumber", "Breaking news, analisis, timestamp, sumber"),
        ("Builder Web3", "Mengikuti teknologi, regulasi, dan proyek", "Konten terlalu berorientasi harga", "Rubrik teknologi, DeFi, keamanan, regulasi"),
        ("Penulis/Editor", "Membuat dan menerbitkan konten berkualitas", "Workflow manual dan revisi tidak terlacak", "Draft, review, schedule, revision history"),
    ]
    add_table(doc, ["Persona", "Tujuan", "Hambatan", "Kebutuhan"], personas, [1700, 2500, 2500, 2660], compact=True)

    add_section(doc, "4. Ruang Lingkup Produk", 1)
    add_section(doc, "4.1 Fitur MVP", 2)
    scope_rows = [
        ("P0", "Pengalaman membaca", "Beranda, listing, detail artikel, kategori, tag, penulis"),
        ("P0", "Editorial", "Dashboard, draft, review, schedule, publish, archive, preview"),
        ("P0", "Kepercayaan", "Sumber, disclosure, tanggal update, koreksi, label konten"),
        ("P0", "SEO", "Metadata, canonical, JSON-LD, sitemap, robots, RSS"),
        ("P0", "Keamanan", "Auth, role, RLS, sanitasi, upload aman, audit log"),
        ("P1", "Discovery", "Pencarian, filter, artikel terkait"),
        ("P1", "Distribusi", "Newsletter opt-in dan unsubscribe"),
        ("P1", "Operasional", "Analytics, error monitoring, health checks"),
    ]
    add_table(doc, ["Prioritas", "Area", "Cakupan"], scope_rows, [1100, 2300, 5960])
    add_section(doc, "4.2 Fase pasca-MVP", 2)
    doc.add_paragraph("Komentar/forum dengan moderasi, bookmark, personalisasi, notifikasi, data pasar real-time, multilingual, dan monetisasi transparan akan dievaluasi setelah MVP memiliki metrik penggunaan dan kemampuan operasional yang memadai.")

    add_section(doc, "5. Information Architecture", 1)
    add_section(doc, "5.1 Sitemap publik", 2)
    for text in [
        "/ — Beranda editorial",
        "/berita, /analisis, /edukasi — rubrik utama",
        "/kategori/[slug] dan /tag/[slug] — arsip tematik",
        "/artikel/[slug] — detail artikel",
        "/penulis/[username] — profil dan artikel penulis",
        "/pencarian — hasil pencarian",
        "/tentang, /kontak, /kebijakan-editorial, /koreksi, /privasi, /disclaimer",
        "/rss.xml, /sitemap.xml, /robots.txt — distribusi dan crawling",
    ]:
        add_bullet(doc, text)
    add_section(doc, "5.2 Taksonomi", 2)
    doc.add_paragraph("Kategori awal: Berita, Analisis, Edukasi, Bitcoin, Ethereum, Altcoin, DeFi, NFT & Web3, Regulasi, dan Keamanan. Jenis konten: berita faktual, analisis, opini, edukasi, siaran pers, dan sponsor.")
    add_callout(doc, "Aturan", "Kategori harus stabil dan kuratorial. Tag dapat lebih granular, tetapi tidak boleh menduplikasi kategori tanpa alasan editorial.", PALE_BLUE, BLUE)

    add_section(doc, "6. Kebutuhan Fungsional", 1)
    requirements = [
        ("FR-01", "Beranda", "Menampilkan headline, breaking news, terbaru, pilihan editor, kategori, dan CTA newsletter.", "P0"),
        ("FR-02", "Artikel", "Menampilkan konten, penulis, timestamp, waktu baca, sumber, disclosure, tag, dan artikel terkait.", "P0"),
        ("FR-03", "Discovery", "Listing, filter, pagination stabil, dan pencarian hanya untuk artikel terbit.", "P1"),
        ("FR-04", "Editorial", "Author membuat draft; editor mereview, menjadwalkan, menerbitkan, dan mengarsipkan.", "P0"),
        ("FR-05", "Revision", "Menyimpan riwayat revisi dan audit aksi sensitif.", "P0"),
        ("FR-06", "Media", "Upload tervalidasi, alt text wajib, dan pengelolaan metadata aset.", "P0"),
        ("FR-07", "SEO", "Metadata unik, canonical, JSON-LD, sitemap, robots, RSS, dan redirect slug.", "P0"),
        ("FR-08", "Newsletter", "Opt-in eksplisit, verifikasi bila dipilih, unsubscribe, dan pencatatan consent.", "P1"),
        ("FR-09", "Koreksi", "Artikel dapat menampilkan catatan koreksi material dan waktu pembaruan.", "P0"),
        ("FR-10", "Preview", "Draft dapat dipreview secara aman tanpa dapat diindeks atau ditebak publik.", "P0"),
    ]
    add_table(doc, ["ID", "Area", "Requirement", "Prioritas"], requirements, [900, 1350, 6210, 900], compact=True)

    add_section(doc, "7. Alur Pengguna Utama", 1)
    add_section(doc, "7.1 Pembaca menemukan artikel", 2)
    for text in ["Membuka beranda atau landing kategori.", "Memilih headline, artikel terbaru, atau menggunakan pencarian.", "Membaca artikel dengan informasi sumber dan waktu publikasi yang jelas.", "Menjelajah artikel terkait atau berlangganan newsletter."]:
        add_number(doc, text)
    add_section(doc, "7.2 Author sampai publikasi", 2)
    for text in ["Author membuat draft, metadata, sumber, gambar, dan disclosure.", "Sistem memvalidasi field wajib dan menyimpan revisi.", "Author mengirim artikel ke status in_review.", "Editor meninjau, meminta revisi, atau menyetujui.", "Editor menerbitkan langsung atau menjadwalkan.", "Sistem melakukan revalidation, memperbarui sitemap/RSS, dan mencatat audit."]:
        add_number(doc, text)
    add_section(doc, "7.3 Koreksi artikel", 2)
    for text in ["Editor membuka artikel terbit dan membuat revisi.", "Perubahan material memerlukan catatan koreksi.", "Sistem memperbarui dateModified dan audit log.", "Pembaca melihat waktu pembaruan dan catatan koreksi secara transparan."]:
        add_number(doc, text)

    add_section(doc, "8. Sistem Editorial", 1)
    add_table(doc, ["Role", "Kemampuan utama", "Batasan"], [
        ("Author", "Membuat/edit draft sendiri, upload media, kirim review", "Tidak menerbitkan atau mengubah role"),
        ("Editor", "Review, revisi, schedule, publish, archive", "Tidak mengelola konfigurasi keamanan"),
        ("Admin", "Seluruh fungsi editorial dan manajemen akses", "Aksi sensitif tercatat pada audit log"),
    ], [1500, 4260, 3600])
    add_section(doc, "8.1 State machine", 2)
    add_callout(doc, "Workflow", "Draft -> In Review -> Scheduled atau Published -> Archived. Perubahan artikel terbit menghasilkan revisi baru dan memperbarui timestamp secara tepat.", PALE_CYAN, CYAN)
    add_section(doc, "8.2 Aturan publikasi", 2)
    for text in ["Judul, slug, ringkasan, isi, penulis, jenis konten, dan tanggal harus valid.", "Sumber wajib untuk klaim faktual yang material.", "Alt text wajib untuk gambar informatif.", "Konten sponsor dan konflik kepentingan wajib memiliki disclosure.", "Analisis harga wajib memuat disclaimer bukan nasihat keuangan.", "AI tidak boleh menerbitkan tanpa review dan persetujuan manusia."]:
        add_bullet(doc, text)

    add_section(doc, "9. Model Data Konseptual", 1)
    entities = [
        ("profiles", "Identitas, username, bio, avatar, role"),
        ("articles", "Konten utama, status, tipe, author/editor, timestamp, SEO"),
        ("article_revisions", "Snapshot perubahan dan actor"),
        ("categories / tags", "Taksonomi dan relasi artikel"),
        ("media", "File, tipe, ukuran, dimensi, alt text, owner"),
        ("article_sources", "URL, label, publisher, tanggal akses"),
        ("newsletter_subscribers", "Email, status consent, timestamps"),
        ("audit_logs", "Actor, aksi, target, metadata aman, timestamp"),
    ]
    add_table(doc, ["Entitas", "Tujuan"], entities, [2800, 6560])
    add_section(doc, "9.1 Aturan data", 2)
    for text in ["Waktu disimpan dalam UTC; editorial default Asia/Jakarta.", "Slug unik dan redirect disimpan ketika slug berubah.", "Invariant penting ditegakkan dengan constraint database.", "RLS aktif untuk tabel yang terekspos.", "Secret, seed phrase, private key, dan credential tidak disimpan sebagai konten.", "Migration kecil, dapat ditinjau, dan memiliki strategi rollback."]:
        add_bullet(doc, text)

    add_section(doc, "10. UX dan Design System", 1)
    doc.add_paragraph("Karakter visual harus modern, kredibel, editorial, dan berorientasi teknologi tanpa menyerupai kasino, skema cepat kaya, atau dashboard trading yang agresif.")
    add_section(doc, "10.1 Prinsip UX", 2)
    for text in ["Mobile-first dengan hierarchy konten yang jelas.", "Tipografi nyaman untuk artikel panjang.", "Kontras minimum WCAG AA dan focus state terlihat.", "Loading, empty, error, unauthorized, dan not-found state dirancang eksplisit.", "Animasi ringan dan menghormati prefers-reduced-motion.", "Warna tidak menjadi satu-satunya pembeda status.", "Tidak ada urgency palsu, dark pattern, atau klaim keuntungan."]:
        add_bullet(doc, text)
    add_section(doc, "10.2 Komponen inti", 2)
    doc.add_paragraph("Header/navigation, breaking-news bar, article card, category chip, author block, source list, disclosure block, newsletter form, search/filter controls, pagination, rich-text renderer, editorial status badge, media picker, dan confirmation dialog.")

    add_section(doc, "11. SEO dan Distribusi", 1)
    seo_rows = [
        ("Metadata", "Title/description unik, canonical, Open Graph, social cards"),
        ("Structured data", "NewsArticle/Article, BreadcrumbList, author, datePublished, dateModified"),
        ("Crawling", "Sitemap canonical; noindex untuk admin, draft, preview, dan pencarian bila diputuskan"),
        ("Distribusi", "RSS dan newsletter"),
        ("URL", "Slug stabil, redirect permanen, status 404 benar"),
        ("Konten", "Tidak ada keyword stuffing, thin pages massal, atau konten AI tersamarkan"),
    ]
    add_table(doc, ["Area", "Requirement"], seo_rows, [2200, 7160])

    add_section(doc, "12. Kredibilitas dan Tata Kelola", 1)
    for text in [
        "Prioritaskan sumber primer dan tampilkan sumber artikel.",
        "Pisahkan fakta, opini, prediksi, dan interpretasi.",
        "Angka pasar menyebut provider dan waktu pengambilan.",
        "Kutipan tidak boleh direkayasa dan harus memiliki atribusi.",
        "Sponsor, afiliasi, kepemilikan aset relevan, dan konflik kepentingan diungkapkan.",
        "Koreksi material terlihat oleh pembaca dan memiliki waktu pembaruan.",
        "Sediakan kebijakan editorial, koreksi, privasi, disclaimer, dan kontak redaksi.",
    ]:
        add_bullet(doc, text)

    add_section(doc, "13. Keamanan dan Privasi", 1)
    security = [
        ("SEC-01", "Least privilege untuk role aplikasi dan database", "P0"),
        ("SEC-02", "RLS deny-by-default dan test untuk setiap role", "P0"),
        ("SEC-03", "Validasi server dan sanitasi rich text berbasis allowlist", "P0"),
        ("SEC-04", "Validasi signature, MIME, ukuran, dan akses upload", "P0"),
        ("SEC-05", "Rate limit login, newsletter, kontak, dan endpoint rentan", "P0"),
        ("SEC-06", "Secret hanya di server; redaction pada log", "P0"),
        ("SEC-07", "CSP dan security headers diuji bertahap", "P1"),
        ("SEC-08", "Consent newsletter, unsubscribe, minimisasi dan retensi data", "P0"),
    ]
    add_table(doc, ["ID", "Requirement", "Prioritas"], security, [1100, 7160, 1100], compact=True)

    add_section(doc, "14. Persyaratan Nonfungsional", 1)
    nfr = [
        ("Performa", "Core Web Vitals kategori baik; JavaScript klien minimal; gambar/font optimal"),
        ("Aksesibilitas", "Target WCAG 2.2 AA; keyboard, semantic HTML, label, focus, error announcement"),
        ("Reliability", "Error handling konsisten, retry/idempotency bila relevan, backup dan recovery"),
        ("Observability", "Structured logs, error monitoring, health check, correlation ID bila diperlukan"),
        ("Maintainability", "TypeScript strict, modular domain, migration dan ADR terdokumentasi"),
        ("Compatibility", "Browser modern; mobile viewport menjadi acceptance path utama"),
    ]
    add_table(doc, ["Area", "Target"], nfr, [2200, 7160])
    add_callout(doc, "Target awal", "Lighthouse Performance, Accessibility, Best Practices, dan SEO >= 90 pada halaman publik representatif dalam kondisi pengujian yang disepakati.", PALE_BLUE, BLUE)

    add_section(doc, "15. Analytics dan Metrik", 1)
    metrics = [
        ("Acquisition", "Organic visits, referral mix, newsletter sign-up rate", "Menilai discovery dan distribusi"),
        ("Engagement", "Returning readers, depth of read, related-article CTR", "Menilai relevansi konten"),
        ("Search", "Search success, zero-result rate", "Menilai findability"),
        ("Editorial", "Draft-to-publish time, revision count, source completeness", "Menilai efisiensi dan kualitas"),
        ("Trust", "Material corrections, disclosure completeness", "Menilai disiplin editorial"),
        ("Quality", "Core Web Vitals, error rate, uptime", "Menilai pengalaman dan reliability"),
    ]
    add_table(doc, ["Kelompok", "Metrik", "Tujuan"], metrics, [1700, 4300, 3360], compact=True)
    doc.add_paragraph("Baseline dan target numerik bisnis ditetapkan setelah analytics aktif dan tersedia data 4-8 minggu. Hindari target vanity metric yang mendorong clickbait.")

    add_section(doc, "16. Roadmap Implementasi", 1)
    roadmap = [
        ("0. Discovery", "Kebutuhan, brand, sitemap, content model, ADR, metrik", "Keputusan inti disetujui"),
        ("1. Fondasi", "Scaffold, CI, design tokens, database, auth, role, RLS", "Build dan security baseline lulus"),
        ("2. Reading", "Beranda, artikel, kategori, pencarian, SEO, RSS", "Alur pembaca lulus E2E"),
        ("3. Editorial", "Dashboard, editor, workflow, media, revision, preview", "Alur author-editor lulus E2E"),
        ("4. Trust", "Newsletter, kebijakan, koreksi, analytics, observability", "Operasional siap staging"),
        ("5. Hardening", "Audit, accessibility, performance, backup, release", "Quality gates produksi lulus"),
    ]
    add_table(doc, ["Fase", "Cakupan", "Exit criteria"], roadmap, [1600, 4700, 3060], compact=True)

    add_section(doc, "17. Acceptance Criteria MVP", 1)
    criteria = [
        ("AC-01", "Pembaca dapat membuka beranda, kategori, dan artikel terbit pada mobile/desktop."),
        ("AC-02", "Artikel menampilkan author, publish/update time, tipe, sumber, disclosure, dan metadata SEO."),
        ("AC-03", "Pengguna tanpa izin tidak dapat mengakses fungsi editorial atau draft."),
        ("AC-04", "Author dapat membuat draft dan mengirim review; editor dapat publish/schedule/archive."),
        ("AC-05", "RLS dan authorization diuji untuk seluruh role dan operasi sensitif."),
        ("AC-06", "Rich text dan upload tervalidasi; script atau file berbahaya ditolak."),
        ("AC-07", "Search/filter memiliki URL stabil, pagination, dan empty state."),
        ("AC-08", "Canonical, JSON-LD, sitemap, robots, RSS, dan redirect slug tervalidasi."),
        ("AC-09", "Newsletter consent dan unsubscribe bekerja serta tercatat."),
        ("AC-10", "Lint, typecheck, unit/integration/E2E test, dan production build lulus."),
        ("AC-11", "Halaman kritis lolos keyboard check dan automated accessibility audit."),
        ("AC-12", "Monitoring, backup, recovery, staging smoke test, dan release checklist tersedia."),
    ]
    add_table(doc, ["ID", "Kriteria"], criteria, [1100, 8260])

    add_section(doc, "18. Risiko dan Mitigasi", 1)
    risks = [
        ("Informasi salah/menyesatkan", "Tinggi", "Workflow review, sumber primer, disclosure, koreksi"),
        ("Akun editor diambil alih", "Tinggi", "MFA, least privilege, session policy, audit log"),
        ("XSS dari rich text", "Tinggi", "Allowlist sanitization, CSP, test payload berbahaya"),
        ("Konten sponsor tidak transparan", "Tinggi", "Label wajib dan approval editorial"),
        ("Vendor lock-in", "Sedang", "Boundary adapter, export data, ADR, dependency review"),
        ("Biaya API/data pasar", "Sedang", "Caching, quota, budget alert, fallback terdefinisi"),
        ("SEO duplicate/thin content", "Sedang", "Canonical, taxonomy governance, content audit"),
        ("Spam newsletter/kontak", "Sedang", "Rate limit, honeypot/CAPTCHA berbasis risiko"),
    ]
    add_table(doc, ["Risiko", "Dampak", "Mitigasi"], risks, [2800, 1100, 5460], compact=True)

    add_section(doc, "19. Dependency dan Keputusan Terbuka", 1)
    add_section(doc, "19.1 Asumsi teknologi", 2)
    doc.add_paragraph("Asumsi awal: Next.js App Router, TypeScript strict, Tailwind CSS, shadcn/ui, PostgreSQL/Supabase, Vercel, Vitest, Playwright, layanan email transaksional, analytics yang menghormati privasi, dan error monitoring. Semua keputusan provider harus dikonfirmasi sebelum provisioning produksi.")
    add_section(doc, "19.2 Pertanyaan terbuka", 2)
    questions = [
        "Apakah MVP hanya berbahasa Indonesia atau harus siap multilingual?",
        "Apakah logo, warna, tipografi, dan brand guideline sudah final?",
        "Siapa role redaksi aktual dan apakah publikasi memerlukan satu atau dua approval?",
        "Apakah editor menggunakan Markdown, rich text, atau block editor?",
        "Provider data pasar, email, analytics, dan monitoring apa yang disetujui?",
        "Apakah sponsor diperbolehkan dan bagaimana workflow komersial-editorial dipisahkan?",
        "Target traffic, anggaran bulanan layanan, SLA, dan tanggal rilis berapa?",
        "Apakah ada konten lama yang harus dimigrasikan?",
        "Apakah komentar/forum masuk roadmap dan siapa yang bertanggung jawab atas moderasi?",
    ]
    for q in questions:
        add_bullet(doc, q)

    add_section(doc, "20. Definition of Done dan Rilis", 1)
    add_section(doc, "20.1 Definition of Done fitur", 2)
    for text in [
        "Acceptance criteria terpenuhi dan alur nyata diverifikasi.",
        "Authorization, validation, loading, empty, error, unauthorized, dan not-found state relevan tersedia.",
        "UI responsif, dapat digunakan dengan keyboard, dan metadata SEO diperbarui.",
        "Test relevan ditambahkan; formatter, lint, typecheck, test, dan build lulus.",
        "Tidak ada secret, error console penting, atau perubahan kontrak tanpa dokumentasi.",
        "Migration, environment example, ADR, dan dokumentasi diperbarui bila relevan.",
    ]:
        add_bullet(doc, text)
    add_section(doc, "20.2 Quality gates produksi", 2)
    for text in [
        "CI dan E2E kritis lulus.",
        "Migration diuji pada staging dan memiliki rollback.",
        "RLS, role, preview draft, sanitasi, upload, dan rate limit diaudit.",
        "Sitemap, canonical, RSS, JSON-LD, accessibility, dan Core Web Vitals diverifikasi.",
        "Backup, recovery, monitoring, alert, dan incident owner tersedia.",
        "Kebijakan editorial, privasi, disclaimer, dan koreksi sesuai implementasi aktual.",
        "Tidak ada placeholder, credential, atau data dummy yang bocor ke produksi.",
    ]:
        add_bullet(doc, text)

    add_section(doc, "Lampiran A — Instruksi Penggunaan PRD untuk Codex", 1)
    add_callout(doc, "Instruksi", "Baca PRD-Crypto-Exist.docx dan AGENTS.md sebelum mengubah kode. Perlakukan scope, prioritas, security baseline, dan Definition of Done sebagai constraint. Jika requirement ambigu atau bertentangan, jelaskan konflik sebelum membuat keputusan permanen.", PALE_CYAN, CYAN)
    doc.add_paragraph("Untuk setiap fitur, Codex harus memetakan requirement ID, menyatakan acceptance criteria, membuat perubahan kecil yang dapat diverifikasi, menambahkan test, menjalankan quality gates yang tersedia, dan melaporkan asumsi serta risiko tersisa.")

    core = doc.core_properties
    core.title = "Crypto Exist — Product Requirements Document"
    core.subject = "PRD platform media berita dan edukasi Web3 serta blockchain"
    core.author = "Crypto Exist"
    core.keywords = "Crypto Exist, PRD, Web3, blockchain, crypto, media"
    core.comments = "Dokumen produk internal versi 1.0"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
